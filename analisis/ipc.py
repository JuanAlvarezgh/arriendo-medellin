"""Serie del indice de precios al consumidor del DANE, componente de arriendo.

El raspado de portales da un corte transversal: dice donde esta caro hoy, no
cuanto subio. La dimension temporal tiene que venir de fuente oficial, y esta
es la razon: los avisos antiguos que siguen publicados son los que nadie tomo,
y una serie construida con sus fechas de publicacion muestra una caida que no
ocurrio.

El DANE publica un anexo mensual en Excel. De cada uno se extraen dos cosas:

  - el indice de la subclase 04130100, arriendo efectivo, a nivel nacional;
  - la variacion anual de la division de alojamiento para Medellin.

La subclase de arriendo solo se publica a nivel nacional. Para Medellin lo mas
fino disponible es la division completa, que incluye servicios publicos. Esa
limitacion se declara en cada salida y no se disimula.
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from datetime import date
from pathlib import Path

import pandas as pd

BASE = "https://www.dane.gov.co/files/operaciones/IPC"
DIR_EXTERNOS = Path("datos/externos/ipc")

MESES = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]

SUBCLASE_ARRIENDO = "04130100"
SUBCLASE_IMPUTADO = "04230100"
DIVISION_ALOJAMIENTO = "alojamiento"

# El anexo con este nombre solo existe desde enero de 2024.
PRIMER_PERIODO = (2024, 1)


def _sin_tildes(texto) -> str:
    return "".join(
        c for c in unicodedata.normalize("NFKD", str(texto)) if not unicodedata.combining(c)
    )


def nombre_periodo(anio: int, mes: int) -> str:
    return f"{MESES[mes - 1]}{anio}"


def url_anexo(anio: int, mes: int) -> str:
    p = nombre_periodo(anio, mes)
    return f"{BASE}/{p}/anex-IPC-{p}.xlsx"


def periodos(hasta: tuple[int, int]) -> list[tuple[int, int]]:
    salida = []
    anio, mes = PRIMER_PERIODO
    while (anio, mes) <= hasta:
        salida.append((anio, mes))
        mes += 1
        if mes > 12:
            mes, anio = 1, anio + 1
    return salida


def descargar(anio: int, mes: int, directorio: Path = DIR_EXTERNOS) -> Path | None:
    """Descarga el anexo de un mes. Devuelve None si el DANE no lo publica."""
    import httpx

    directorio.mkdir(parents=True, exist_ok=True)
    destino = directorio / f"anex-IPC-{nombre_periodo(anio, mes)}.xlsx"
    if destino.exists() and destino.stat().st_size > 10_000:
        return destino

    try:
        r = httpx.get(
            url_anexo(anio, mes),
            timeout=90.0,
            follow_redirects=True,
            headers={"User-Agent": "rastreador-academico/0.1 (analitica de datos)"},
        )
    except httpx.HTTPError:
        return None
    if r.status_code != 200 or len(r.content) < 10_000:
        return None
    destino.write_bytes(r.content)
    return destino


@dataclass
class Observacion:
    periodo: date
    indice_arriendo: float | None = None
    variacion_mensual: float | None = None
    variacion_anual: float | None = None
    indice_imputado: float | None = None
    alojamiento_medellin_anual: float | None = None
    alojamiento_nacional_anual: float | None = None


def _leer_subclase(libro, codigo: str) -> tuple[float | None, float | None, float | None]:
    """Devuelve (indice, variacion mensual, variacion anual) de una subclase."""
    if "8" not in libro.sheetnames:
        return None, None, None
    hoja = libro["8"]
    for fila in hoja.iter_rows(values_only=True):
        if fila and str(fila[0]).strip() == codigo:
            def num(v):
                try:
                    return float(v)
                except (TypeError, ValueError):
                    return None

            return num(fila[2]), num(fila[3]), num(fila[5])
    return None, None, None


def _leer_ciudad(libro, ciudad: str) -> tuple[float | None, float | None]:
    """Variacion anual de la division de alojamiento: (ciudad, nacional)."""
    if "6" not in libro.sheetnames:
        return None, None
    hoja = libro["6"]

    columna = None
    valor_ciudad = valor_total = None
    objetivo = _sin_tildes(ciudad).lower()

    for fila in hoja.iter_rows(values_only=True):
        if not fila or fila[0] is None:
            continue
        etiqueta = _sin_tildes(fila[0]).strip().lower()

        if columna is None and etiqueta.startswith("ciudad"):
            for i, celda in enumerate(fila):
                if celda and DIVISION_ALOJAMIENTO in _sin_tildes(celda).lower():
                    columna = i
                    break
            continue

        if columna is None:
            continue
        try:
            valor = float(fila[columna])
        except (TypeError, ValueError, IndexError):
            continue
        if etiqueta.startswith("total"):
            valor_total = valor
        elif etiqueta.startswith(objetivo):
            valor_ciudad = valor

    return valor_ciudad, valor_total


def leer_anexo(ruta: Path, anio: int, mes: int) -> Observacion:
    import openpyxl

    libro = openpyxl.load_workbook(ruta, data_only=True, read_only=False)
    indice, mensual, anual = _leer_subclase(libro, SUBCLASE_ARRIENDO)
    imputado, _, _ = _leer_subclase(libro, SUBCLASE_IMPUTADO)
    medellin, nacional = _leer_ciudad(libro, "medellin")
    libro.close()

    return Observacion(
        periodo=date(anio, mes, 1),
        indice_arriendo=indice,
        variacion_mensual=mensual,
        variacion_anual=anual,
        indice_imputado=imputado,
        alojamiento_medellin_anual=medellin,
        alojamiento_nacional_anual=nacional,
    )


def construir_serie(hasta: tuple[int, int] = (2026, 7), verbose: bool = True) -> pd.DataFrame:
    """Descarga y arma la serie completa. Los meses no publicados se omiten."""
    filas: list[Observacion] = []
    for anio, mes in periodos(hasta):
        ruta = descargar(anio, mes)
        if ruta is None:
            if verbose:
                print(f"  {nombre_periodo(anio, mes)}: no publicado")
            continue
        try:
            obs = leer_anexo(ruta, anio, mes)
        except Exception as e:
            if verbose:
                print(f"  {nombre_periodo(anio, mes)}: ilegible ({type(e).__name__})")
            continue
        filas.append(obs)
        if verbose:
            print(
                f"  {nombre_periodo(anio, mes)}: indice {obs.indice_arriendo}  "
                f"anual {obs.variacion_anual}%  Medellin alojamiento {obs.alojamiento_medellin_anual}%"
            )

    df = pd.DataFrame([f.__dict__ for f in filas])
    if not df.empty:
        df = df.sort_values("periodo").reset_index(drop=True)
    return df


def acumulado(df: pd.DataFrame, columna: str = "indice_arriendo") -> float | None:
    """Crecimiento porcentual entre el primer y el ultimo indice de la serie."""
    s = df[columna].dropna()
    if len(s) < 2:
        return None
    return (s.iloc[-1] / s.iloc[0] - 1) * 100
